"""Builds helpers/caRefreshData.json for the "CA Refresh" project from
  - the CA Refresh test-case workbook (Pre-condition column + Coverage Lists sheet)
  - the Donegal Commercial Auto product model (Coverages + Coverage Term Options sheets)

  python -I scripts/buildCaRefreshData.py "<test cases.xlsx>" "<product model.xlsx>"
Defaults to the V:\\Amit Int\\CARefresh copies. Offline; re-run whenever either workbook changes.
"""
import json, os, re, sys, warnings
import openpyxl

warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = r'V:\Amit Int\CARefresh'
TC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(SRC, 'ABP-1775_CA Refresh Docs - Batch 2.xlsx')
PM = sys.argv[2] if len(sys.argv) > 2 else os.path.join(SRC, 'Donegal Commercial Auto Product Model _SBT_OCP_v1.0.xlsx')
OUT = os.path.join(HERE, '..', 'helpers', 'caRefreshData.json')


from openpyxl.cell.rich_text import CellRichText, TextBlock


def effective(c):
    """A cell's value without struck-through text; a fully struck cell is no longer valid (empty)."""
    v = c.value
    if v is None:
        return None
    if getattr(c.font, 'strike', False):
        return ''
    if isinstance(v, CellRichText):
        return ''.join(p if isinstance(p, str) else (p.text if not p.font.strike else '') for p in v).strip()
    return v


def norm(s):
    return re.sub(r'[^a-z0-9]+', ' ', str(s or '').lower()).strip()


def clean(s):
    return re.sub(r'\s+', ' ', str(s or '').replace('\ufffd', '-')).strip()


# ── Test cases: id, LOB, states, precondition ────────────────────────────────
wb = openpyxl.load_workbook(TC, read_only=True)
ws = wb['CA Refresh Doc Upd - Batch1']
cases = []
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    m = re.search(r'For States?:\s*([A-Z, ]+?),\s*Verify', str(r[1]))
    states = [s.strip() for s in m.group(1).split(',') if s.strip()] if m else []
    lob = 'AUTO_DEALERS' if 'Auto Dealers Policy' in str(r[1]) else 'BUSINESS_AUTO'
    cases.append({'id': str(r[0]), 'lob': lob, 'states': states, 'name': clean(r[1]), 'pre': str(r[3] or '')})


def parse_pre(pre):
    """Section -> [coverage names], plus the 'Additional Coverages from mapping' (forms) list."""
    sections, forms, cur, mapping = [], [], None, False
    for line in pre.splitlines():
        t = line.strip()
        if not t:
            continue
        if t.startswith('Coverages:'):
            continue
        if t.startswith('Additional Coverages from mapping'):
            mapping, cur = True, None
            continue
        if t.endswith(':-'):
            mapping = False
            cur = {'category': t[:-2].strip(), 'coverages': []}
            sections.append(cur)
            continue
        if mapping:
            forms.append(clean(t))
        elif cur is not None and t not in cur['coverages']:
            cur['coverages'].append(clean(t))
    return sections, forms


# ── Product model lookups ────────────────────────────────────────────────────
pm = openpyxl.load_workbook(PM, read_only=True)
pmr = openpyxl.load_workbook(PM, read_only=True, rich_text=True)


def wb_overrides(sheet, col, maxr):
    """row number -> availability text with struck-through parts removed (col is 1-based)."""
    return {row[0].row: effective(row[0]) for row in pmr[sheet].iter_rows(min_row=3, max_row=maxr, min_col=col, max_col=col) if row[0].value is not None}


def apply_wb(rows, over, idx):
    out = []
    for k, r in enumerate(rows):
        if (k + 1) in over:
            r = list(r)
            r[idx] = over[k + 1]
        out.append(r)
    return out


cov_rows = apply_wb(list(pm['Coverages'].iter_rows(values_only=True, max_col=40)), wb_overrides('Coverages', 11, 600), 10)
cov_rows = apply_wb(cov_rows, wb_overrides('Coverages', 5, 600), 4)  # coverage name without struck-through (old) names
hdr = cov_rows[1]
STATE_COLS = {hdr[i]: i for i in range(15, 39) if hdr[i]}
cov_index = {}
for r in cov_rows[2:]:
    if not r[4]:
        continue
    key = (norm(r[3]), norm(r[4]))
    cov_index[key] = r
    cov_index.setdefault((None, norm(r[4])), r)

TARGET_STATES = {st for c in cases for st in c['states']}
opt_rows = apply_wb(list(pm['Coverage Term Options'].iter_rows(values_only=True, max_col=50)), wb_overrides('Coverage Term Options', 14, 6000), 13)
ohdr = opt_rows[1]
OPT_STATE_COLS = {ohdr[i]: i for i in range(15, 39) if ohdr[i]}
terms = {}  # (category, coverage) -> {term: [{"option", "code", "states": [...], "wb": "Y/N", "default": bool}]}
# State cells can hold two lines - the old value above the new one ("Y" / "N"), the old one sometimes struck through.
# QA (Devyani, 10/30): the CURRENT value is the last line that is not struck through - "Y<newline>N" = no longer offered,
# "N<newline>Y" = offered. Read those cells with rich text so struck parts are dropped.
def state_now(cell):
    v = effective(cell)
    lines = [x.strip().upper() for x in str(v or '').split(chr(10)) if x.strip()]
    return lines[-1] if lines else ''
opt_state_now = {}  # sheet row -> {state: current value}
for rn, row in enumerate(pmr['Coverage Term Options'].iter_rows(min_row=3, max_row=6000, min_col=1, max_col=39), start=3):
    if len(row) < 39 or (row[0].value is None and row[4].value is None):
        continue
    opt_state_now[rn] = {st: state_now(row[c]) for st, c in OPT_STATE_COLS.items() if st in TARGET_STATES and row[c].value is not None}
for k, r in enumerate(opt_rows[2:], start=3):
    if not r[4] or not r[6] or not r[9]:
        continue
    # QA (Devyani, 10/30): every term and option of an in-scope coverage is checked - no 'SBT Update 2025' filter here
    # (Hired Auto Collision: only Deductible was SBT-flagged, so Coverage Type and its dependent deductibles were missed).
    t = terms.setdefault((norm(r[3]), norm(r[4])), {}).setdefault(clean(r[6]), [])
    states = [st for st, v in opt_state_now.get(k, {}).items() if v == 'Y']
    t.append({'option': clean(r[9]), 'code': clean(r[8]) or None, 'value': clean(r[10]) or None,
              'default': str(r[12] or '').strip().upper().startswith('Y'), 'wb': re.sub(r'\s+', '', str(r[13] or '')).upper() or None, 'states': states})


# Conditions + Exclusions rows in the same tuple shape as the Coverages rows, so one lookup covers all three sheets
# (the precondition also lists exclusions such as War Exclusion, which are not on the Coverages tab).
def _synthetic(sheet, wb_i, form_i, state0, maxr=4000):
    rows = apply_wb(list(pm[sheet].iter_rows(values_only=True, max_col=60)), wb_overrides(sheet, wb_i + 1, maxr), wb_i)
    hdr = rows[1]
    scols = {hdr[k]: k for k in range(state0, state0 + 24) if hdr[k]}
    out = []
    for r in rows[2:]:
        if not r[4]:
            continue
        row = [None] * 40
        row[3], row[4], row[5], row[10], row[12] = r[3], r[4], r[5], r[wb_i], r[form_i]
        for st, k in scols.items():
            if st in STATE_COLS:
                row[STATE_COLS[st]] = r[k]
        out.append(row)
    return out


import re as _re
_all_rows = [r for r in cov_rows[2:] if r[4]] + _synthetic('Conditions', 9, 11, 14) + _synthetic('Exclusions', 9, 11, 14)
_by_name = {}
_by_form = {}
for _r in _all_rows:
    for _nm in str(_r[4]).split(chr(10)):
        if _nm.strip():
            _by_name.setdefault(norm(_nm), []).append(_r)
    _fm = _re.search(r'\b(CAD?)\s?(\d{2})\s?(\d{2})\b', str(_r[12] or '') + ' ' + str(_r[4]), _re.I)
    if _fm:
        _by_form.setdefault((_fm.group(1) + _fm.group(2) + _fm.group(3)).upper(), []).append(_r)


def canon(name):
    n = norm(name)
    n = _re.sub(r'^auto dealers? ', 'garage ', n)          # precondition says "Auto Dealers X", the model says "Garage X"
    return _re.sub(r'^(cad?) ?\d{2} ?\d{2} ', '', n)         # drop a leading form number


_by_canon = {}
for _r in _all_rows:
    for _nm in str(_r[4]).split(chr(10)):
        if _nm.strip():
            _by_canon.setdefault(canon(_nm), []).append(_r)


def _jacc(a, b):
    ta, tb = set(a.split()), set(b.split())
    return len(ta & tb) / len(ta | tb) if ta | tb else 0


def _pick(rows, category):
    """Among several rows for the same coverage (one per effective date / category) prefer the same category, then any WriteBiz=Y."""
    same = [r for r in rows if norm(r[3]) == norm(category)] or rows
    return next((r for r in same if r[10] and 'y' in str(r[10]).lower()), same[0])


def find_model_row(category, name):
    rows = _by_name.get(norm(name)) or _by_canon.get(canon(name))
    if not rows:
        fm = _re.search(r'\b(CAD?)\s?(\d{2})\s?(\d{2})\b', name, _re.I)
        if fm:
            rows = _by_form.get((fm.group(1) + fm.group(2) + fm.group(3)).upper())
    if not rows:
        c = canon(name)
        best = max(_by_canon, key=lambda k: _jacc(c, k), default=None)
        if best and _jacc(c, best) >= 0.8:
            rows = _by_canon[best]
    return _pick(rows, category) if rows else None


def enrich(category, name):
    row = find_model_row(category, name)
    out = {'name': name}
    if row:
        out['code'] = row[5]
        out['form'] = clean(row[12]) or None
        out['writeBiz'] = clean(row[10]) or None  # Donegal WriteBiz Availability For Agents (Y/N), struck-through text removed
        out['availability'] = {s: clean(row[i]) for s, i in STATE_COLS.items() if row[i] not in (None, '')}
    t = terms.get((norm(category), norm(name)))
    if not t and row:
        t = terms.get((norm(row[3]), norm(row[4])))
    if t:
        out['terms'] = t
    return out


# ── Coverage Lists sheet: the "Removed" column ───────────────────────────────
wb2 = openpyxl.load_workbook(TC, read_only=True)
removed, cur = [], None
for r in wb2['Coverage Lists'].iter_rows(min_row=4, values_only=True):
    t = clean(r[5]) if r[5] else ''
    if not t:
        continue
    if t.endswith(':-'):
        cur = t[:-2].strip()
    else:
        removed.append({'category': cur, 'name': t})

# Every product-model coverage that has term options (not only the precondition list): lets the run compare the options
# of ANY coverage card WriteBiz shows against the model.
model_terms = []
for (ncat, ncov), tmap in terms.items():
    row = find_model_row(ncat, ncov) or cov_index.get((ncat, ncov)) or cov_index.get((None, ncov))
    raw = str(row[4]) if row else ncov
    for nm in [x for x in raw.split(chr(10)) if x.strip()] or [ncov]:
        model_terms.append({'category': clean(row[3]) if row else ncat, 'name': clean(nm), 'terms': tmap})

# Every term name per coverage (Coverage Terms tab; includes free-text terms that have no option rows) so a dialog's terms can be
# checked for presence. The tab is declared with ~1M blank rows, so only the first few thousand are read.
term_names = {}
for r in pm['Coverage Terms'].iter_rows(min_row=2, max_row=6000, values_only=True, max_col=60):
    if r[4] and r[6]:  # every term (QA 10/30: no 'SBT Update 2025' filter on the Coverage Terms tab)
        for nm in str(r[4]).split(chr(10)):
            if nm.strip():
                lst = term_names.setdefault(norm(nm), [])
                if clean(r[6]) not in lst:
                    lst.append(clean(r[6]))

# ── "SBT Update" = yes on the Coverages / Conditions / Exclusions tabs: the full CA Refresh set (the test-case
#    precondition only lists some of it) ───────────────────────────────────────────────────────────────────────────
def sbt_rows(sheet, kind, name_i, cat_i, code_i, wb_i, form_i, state0, sbt_i, maxc):
    rows = apply_wb(list(pm[sheet].iter_rows(values_only=True, max_col=maxc)), wb_overrides(sheet, wb_i + 1, 4000), wb_i)
    # Renamed coverages keep the old name struck through in the same cell ("Garage Collision" / "Auto Dealers Collision"):
    # read the name without the struck part, else the joined old+new name is reported MISSING (10/30 QA screenshot).
    rows = apply_wb(rows, wb_overrides(sheet, name_i + 1, 4000), name_i)
    hdr = rows[1]
    scols = {hdr[i]: i for i in range(state0, state0 + 24) if hdr[i] in TARGET_STATES}
    out = []
    for r in rows[2:]:
        if len(r) <= sbt_i or not r[sbt_i] or 'yes' not in str(r[sbt_i]).lower() or not r[name_i]:
            continue
        out.append({'kind': kind, 'category': clean(r[cat_i]), 'name': clean(r[name_i]), 'code': clean(r[code_i]) or None,
                    'form': clean(r[form_i]) or None, 'writeBiz': clean(r[wb_i]) or None,
                    'availability': {st: clean(r[i]) for st, i in scols.items() if r[i] not in (None, '')}})
    return out

sbt = (sbt_rows('Coverages', 'coverage', 4, 3, 5, 10, 12, 15, 48, 50)
       + sbt_rows('Conditions', 'condition', 4, 3, 5, 9, 11, 14, 44, 50)
       + sbt_rows('Exclusions', 'exclusion', 4, 3, 5, 9, 11, 14, 44, 50))

# Clause Availability Rules 2025: which policy types (CA7LinePolicyType = ...) each clause code is offered on
clause = {}
cur = None
for r in pm['Clause Availability Rules 2025'].iter_rows(min_row=5, values_only=True, max_col=8):
    if r[0]:
        cur = clean(r[0])
    if cur and r[4] and str(r[4]).endswith('/CA7LinePolicyType') and r[6]:
        c = clause.setdefault(cur, {'allow': [], 'deny': []})
        (c['deny'] if str(r[5]).strip() in ('!=', '<>') else c['allow']).append(clean(r[6]))
codes = {c['code'] for c in sbt if c['code']} | {c.get('code') for lob in ('BUSINESS_AUTO', 'AUTO_DEALERS')
         for sec in (parse_pre([c0 for c0 in cases if c0['lob'] == lob][0]['pre'])[0] if [c0 for c0 in cases if c0['lob'] == lob] else [])
         for c in [enrich(sec['category'], n) for n in sec['coverages']]}
clause = {k: v for k, v in clause.items() if k in codes}

data = {
    'termNames': term_names,
    'sbt': sbt,
    'clauseRules': clause,
    'modelTerms': model_terms,
    'source': {'testCases': os.path.basename(TC), 'productModel': os.path.basename(PM)},
    'effective': {'IN': '2027-05-01', 'IL': '2027-05-01', 'default': '2027-08-01'},
    'lobs': {},
    'removed': removed,
}
for lob in ('BUSINESS_AUTO', 'AUTO_DEALERS'):
    mine = [c for c in cases if c['lob'] == lob]
    if not mine:
        continue
    sections, forms = parse_pre(mine[0]['pre'])
    identical = all(parse_pre(c['pre']) == (sections, forms) for c in mine)
    data['lobs'][lob] = {
        'testCases': {s: c['id'] for c in mine for s in c['states']},
        'testCaseNames': {c['id']: c['name'] for c in mine},
        'preconditionIdenticalAcrossStates': identical,
        'sections': [{'category': s['category'], 'coverages': [enrich(s['category'], n) for n in s['coverages']]} for s in sections],
        'mappingForms': forms,
    }
data['states'] = sorted({s for c in cases for s in c['states']})

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(data, f, indent=1, ensure_ascii=False)
for lob, d in data['lobs'].items():
    n = sum(len(s['coverages']) for s in d['sections'])
    hit = sum(1 for s in d['sections'] for c in s['coverages'] if c.get('code'))
    print(f"{lob}: {len(d['testCases'])} states, {n} coverages ({hit} matched in product model), {len(d['mappingForms'])} forms, identical={d['preconditionIdenticalAcrossStates']}")
print('states:', ' '.join(data['states']), '| removed:', len(removed), '| SBT rows:', len(sbt), '| clause rules:', len(clause))
